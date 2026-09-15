"""Synthetic in-memory cell objects test parser behavior without authoring XLSX."""

from datetime import date
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

from pipeline.clean import clean_sheet, coordinates, count, evaluate_sum, run, survey_date


class Cells:
    title = 'mock'
    max_row = 7
    max_column = 14

    def __init__(self):
        self.values = {'A1': 'ปริมาณจราจร ประจำเดือนมกราคม 2567',
                       'A3': 'ลำดับที่', 'C3': 'ถนน/ซอย', 'D3': 'ช่วงเวลา',
                       'K4': 'แต่ละช่วงเวลา', 'L4': 'แต่ละถนน', 'M4': 'รวมทั้งแยก',
                       'A5': 1, 'B5': 'สถานที่จำลอง\n(4ม.ค. 67)', 'C6': 'ถนนจำลอง',
                       'N5': '13.75, 100.5', 'L6': '=SUM(K5:K7)', 'M5': '=SUM(L5:L7)'}
        for column, name in zip('EFGHIJ', ['รถยนต์นั่ง', 'ตู้/ปิคอัพ', 'เมล์ใหญ่', 'เมล์เล็ก', 'บรรทุก', 'สามล้อ']):
            self.values[f'{column}4'] = name
        for row, window in zip(range(5, 8), ['7.00 - 9.00', '9.00 - 16.00', '16.00 - 19.00']):
            self.values[f'D{row}'] = f'ช่วงจำลอง ({window} น.)'
            self.values[f'K{row}'] = f'=SUM(E{row}:J{row})'
            for column in 'EFGHIJ':
                self.values[f'{column}{row}'] = 0 if column == 'J' else 10

    def __getitem__(self, address):
        row, _ = coordinate_to_tuple(address)
        return SimpleNamespace(value=self.values.get(address), row=row, coordinate=address)

    def cell(self, row, column):
        return self[f'{get_column_letter(column)}{row}']

    def iter_rows(self, min_row, max_row, min_col, max_col):
        return [[self.cell(r, c) for c in range(min_col, max_col + 1)] for r in range(min_row, max_row + 1)]


def parse(sheet):
    cached = Cells()
    for address in list(cached.values):
        if str(cached.values[address]).startswith('='):
            cached.values[address] = evaluate_sum(cached, address)
    return clean_sheet(sheet, cached, 'synthetic-source')


class Book:
    def __init__(self, sheet):
        self.sheet = sheet

    def __iter__(self):
        return iter([self.sheet])

    def __getitem__(self, name):
        return self.sheet

    def close(self):
        pass


def mock_loader(broken=False):
    def load(path, data_only=False, **kwargs):
        sheet = Cells()
        if data_only:
            for address in list(sheet.values):
                if str(sheet.values[address]).startswith('='):
                    sheet.values[address] = evaluate_sum(sheet, address)
        elif broken:
            sheet.values['E5'] = '=[1]Analysis!$B$97'
        return Book(sheet)
    return load


class CleanTests(unittest.TestCase):
    def test_2022_report_is_excluded_without_changing_source(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'raw'
            source.mkdir()
            raw = source / 'old.xlsx'
            raw.write_bytes(b'Synthetic placeholder; reading mocked')
            loader = mock_loader()
            def old_report(*args, **kwargs):
                book = loader(*args, **kwargs)
                book.sheet.values['A1'] = 'ปริมาณจราจร ประจำเดือนกรกฎาคม 2565'
                return book
            with patch('pipeline.clean.load_workbook', side_effect=old_report):
                result = run(source, Path(directory) / 'output')
            self.assertEqual(result['source_files'], 0)
            self.assertEqual(result['observations_accepted'], 0)
            self.assertEqual(len(result['excluded_reports']), 1)
            self.assertEqual(raw.read_bytes(), b'Synthetic placeholder; reading mocked')

    def test_entire_survey_quarantined_and_raw_input_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'raw'
            source.mkdir()
            file = source / 'synthetic.xlsx'
            file.write_bytes(b'Synthetic placeholder; workbook reading is mocked')
            output = Path(directory) / 'run'
            before = file.read_bytes()
            with patch('pipeline.clean.load_workbook', side_effect=mock_loader(broken=True)):
                summary = run(source, output)
            self.assertEqual(summary['surveys_quarantined'], 1)
            self.assertEqual(summary['observations_accepted'], 0)
            self.assertEqual(summary['observation_rows_quarantined'], 3)
            self.assertEqual(json.loads((output / 'observations.json').read_text()), [])
            self.assertEqual(len(json.loads((output / 'rejected_rows.json').read_text(encoding='utf-8'))), 3)
            self.assertEqual(file.read_bytes(), before)
            with self.assertRaises(ValueError):
                run(source, output)

    def test_duplicate_candidates_quarantine_both_copies(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'raw'
            source.mkdir()
            for name in ['copy1.xlsx', 'copy2.xlsx']:
                (source / name).write_bytes(b'Synthetic placeholder; reading mocked')
            with patch('pipeline.clean.load_workbook', side_effect=mock_loader()):
                summary = run(source, Path(directory) / 'run')
            self.assertEqual(summary['surveys_quarantined'], 2)
            self.assertEqual(summary['issues_by_code']['duplicate_candidate'], 6)
            self.assertEqual(summary['observations_accepted'], 0)

    def test_multiline_name_date_and_centered_road_label(self):
        surveys, roads, observations, issues, raw = parse(Cells())
        self.assertEqual(surveys[0]['survey_date'], '2024-01-04')
        self.assertEqual(surveys[0]['intersection_name'], 'สถานที่จำลอง')
        self.assertEqual(roads[0]['source_start_row'], 5)
        self.assertEqual(roads[0]['road_name_source_row'], 6)
        self.assertEqual([o['duration_minutes'] for o in observations], [120, 420, 180])
        self.assertEqual(sum(o['vehicle_total'] for o in observations), 150)
        self.assertEqual(issues, [])
        self.assertEqual(raw[0]['cells'][9], 0)

    def test_external_formula_is_not_replaced_by_cached_number(self):
        sheet = Cells()
        sheet.values['E5'] = '=[1]Analysis!$B$97'
        _, _, observations, issues, raw = parse(sheet)
        self.assertTrue(any(i['code'] == 'invalid_observation' for i in issues))
        self.assertEqual(len(observations), 2)
        self.assertEqual(raw[0]['cells'][4], '=[1]Analysis!$B$97')
        self.assertEqual(raw[0]['cached_cells'][4], 10)

    def test_wrong_subtotal_formula_is_rejected(self):
        sheet = Cells()
        sheet.values['L6'] = '=SUM(K5:K6)'
        issues = parse(sheet)[3]
        self.assertTrue(any(i['code'] == 'subtotal_mismatch_or_unverifiable' for i in issues))

    def test_unlabeled_blank_periods_are_not_proven_unused_by_zero_subtotals(self):
        sheet = Cells()
        sheet.max_row = 10
        for r, source_row in zip(range(8, 11), range(5, 8)):
            sheet.values[f'D{r}'] = sheet.values[f'D{source_row}']
            sheet.values[f'K{r}'] = f'=SUM(E{r}:J{r})'
        sheet.values['L9'] = '=SUM(K8:K10)'
        sheet.values['M5'] = '=SUM(L5:L10)'
        self.assertEqual(evaluate_sum(sheet, 'L9'), 0)
        self.assertEqual(evaluate_sum(sheet, 'M5'), 150)
        issues = parse(sheet)[3]
        self.assertTrue(any(i['code'] == 'missing_or_ambiguous_road' and i['severity'] == 'error'
                            for i in issues))

    def test_counts_distinguish_zero_missing_negative_fraction(self):
        self.assertEqual(count(0), 0)
        for value in [None, -1, 1.5, float('nan'), True, '12', '=SUM(A1)']:
            with self.subTest(value=value), self.assertRaises(ValueError):
                count(value)

    def test_date_validation_and_buddhist_leap_year(self):
        self.assertEqual(survey_date('(29 ก.พ. 67)', date(2024, 2, 1)), date(2024, 2, 29))
        self.assertEqual(survey_date('4ม.ค.67)', date(2024, 1, 1)), date(2024, 1, 4))
        for raw, month in [('(29 ก.พ. 68)', date(2025, 2, 1)), ('(1 มี.ค. 67)', date(2024, 2, 1))]:
            with self.assertRaises(ValueError):
                survey_date(raw, month)

    def test_number_in_location_name_is_not_a_date(self):
        sheet = Cells()
        sheet.values['B5'] = 'สถานที่จำลอง (2)\n4 ม.ค.67)'
        surveys, _, _, issues, _ = parse(sheet)
        self.assertEqual(surveys[0]['intersection_name'], 'สถานที่จำลอง (2)')
        self.assertEqual(surveys[0]['survey_date'], '2024-01-04')
        self.assertEqual(issues, [])

    def test_invalid_coordinates_are_not_guessed(self):
        self.assertEqual(coordinates('13.75  100.5'), (13.75, 100.5))
        for value in ['13.7 10.0.609137', '100.5 13.7', '13.7, 100.5 extra']:
            with self.assertRaises(ValueError):
                coordinates(value)

    def test_period_overlap_and_unknown_header(self):
        sheet = Cells()
        sheet.values['D6'] = 'ช่วงจำลอง (8.00 - 16.00 น.)'
        self.assertTrue(any(i['code'] == 'incomplete_or_overlapping_periods' for i in parse(sheet)[3]))
        sheet.values['E4'] = 'unexpected new category'
        with self.assertRaises(ValueError):
            parse(sheet)

    def test_formula_cycles_and_external_links_are_rejected(self):
        sheet = Cells()
        for value in ['=SUM(K5)', '=[1]Analysis!A1', '=1+2']:
            sheet.values['K5'] = value
            with self.assertRaises(ValueError):
                evaluate_sum(sheet, 'K5')


if __name__ == '__main__':
    unittest.main()
