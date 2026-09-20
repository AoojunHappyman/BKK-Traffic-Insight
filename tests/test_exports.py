import csv
from datetime import datetime
from io import BytesIO, StringIO
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from app import create_app
from app.exports import COLUMNS
from app.map_quality import PENDING_COORDINATES


def observation(**changes):
    row = dict(survey_id='a' * 64, observation_id='b' * 64, survey_date='2025-01-02',
               intersection_name='ทดสอบ, สถานที่', road_name='ถนนทดสอบ', period_start='07:00',
               period_end='09:00', duration_minutes=120, passenger_car=10, van_pickup=2,
               large_bus=0, small_bus=1, truck=3, three_wheeler=0, vehicle_total=16,
               latitude=13.75, longitude=100.5, filename='source.xlsx', sheet_name='มค68', source_row=5)
    return {**row, **changes}


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.client = create_app().test_client()

    def get(self, url, rows):
        with patch('app.routes.traffic.query', return_value=rows):
            response = self.client.get(url)
        self.assertEqual(response.status_code, 200)
        return response

    def test_csv_all_rows_thai_bom_quoting_and_numbers(self):
        response = self.get('/api/export/overview.csv', [observation()] * 501)
        self.assertTrue(response.data.startswith(b'\xef\xbb\xbf'))
        rows = list(csv.reader(StringIO(response.data.decode('utf-8-sig'))))
        self.assertEqual(len(rows), 502)
        self.assertEqual(rows[1][1], 'ทดสอบ, สถานที่')
        self.assertEqual(rows[1][12], '16')
        self.assertIn('attachment;', response.headers['Content-Disposition'])
        self.assertEqual(response.headers['X-Export-Rows'], '501')
        self.assertEqual(response.headers['Cache-Control'], 'no-store')

    def test_excel_typed_cells_and_filters_and_formula_safety(self):
        response = self.get('/api/export/vehicles.xlsx?period=07:00-09:00&intersection_name=%3D1%2B1',
                            [observation(intersection_name='=1+1')])
        workbook = load_workbook(BytesIO(response.data), data_only=False)
        sheet = workbook['Traffic data']
        self.assertEqual(sheet['B2'].value, '=1+1')
        self.assertEqual(sheet['B2'].data_type, 's')
        self.assertIsInstance(sheet['A2'].value, datetime)
        self.assertEqual(sheet['M2'].value, 16)
        self.assertEqual(sheet['M2'].data_type, 'n')
        self.assertEqual(sheet['J2'].value, 1)
        info = dict(workbook['Export info'].values)
        self.assertEqual(info['สถานที่'], '=1+1')
        self.assertEqual(info['ช่วงเวลา'], '07:00-09:00')
        self.assertEqual(info['จำนวนรถรวม (คัน)'], 16)
        self.assertEqual(sheet.freeze_panes, 'D2')
        self.assertEqual(sheet.auto_filter.ref, 'A1:T2')
        self.assertFalse(any(c.data_type == 'f' for s in workbook for row in s for c in row))
        workbook.close()

    def test_csv_neutralizes_formula_strings_and_keeps_zero(self):
        for name in ['=1+1', '+1', '-1', '@SUM(A1)', '\t=1+1', '  =1+1']:
            response = self.get('/api/export/overview.csv', [observation(intersection_name=name)])
            rows = list(csv.reader(StringIO(response.data.decode('utf-8-sig'))))
            self.assertEqual(rows[1][1], "'" + name)
            self.assertEqual(rows[1][8], '0')

    def test_export_filters_are_parameterized_and_not_paginated(self):
        selection = {'start_date': '2025-01-01', 'end_date': '2025-01-31',
                     'intersection_name': "x' OR 1=1 --", 'period': '07:00-09:00'}
        with patch('app.routes.traffic.query', return_value=[]) as query:
            self.assertEqual(self.client.get('/api/export/vehicles.csv', query_string=selection).status_code, 200)
        sql, params = query.call_args.args
        self.assertNotIn(selection['intersection_name'], sql)
        self.assertNotIn('LIMIT', sql)
        self.assertEqual(params, ['2025-01-01', '2025-01-31', selection['intersection_name'], '07:00', '09:00'])

    def test_invalid_exports_are_rejected_before_query(self):
        urls = ['overview.pdf', 'other.csv', 'overview.csv?limit=1', 'overview.csv?offset=0',
                'overview.csv?start_date=2025-02-30', 'overview.csv?start_date=2026-01-01&end_date=2025-01-01',
                'overview.csv?period=07:00-09:00', 'vehicles.csv?period=09:00-07:00',
                'vehicles.csv?period=07:00-09:00&period=09:00-17:00', 'overview.csv?unexpected=true']
        with patch('app.routes.traffic.query') as query:
            for url in urls:
                with self.subTest(url=url):
                    self.assertEqual(self.client.get('/api/export/' + url).status_code, 400)
            query.assert_not_called()

    def test_map_matches_coordinate_exclusions_other_pages_keep_counts(self):
        sid, lat, lon = next(iter(PENDING_COORDINATES))
        rows = [observation(), observation(latitude=None, longitude=None),
                observation(survey_id=sid, latitude=lat, longitude=lon), observation(latitude=100)]
        self.assertEqual(self.get('/api/export/map.csv', rows).headers['X-Export-Rows'], '1')
        self.assertEqual(self.get('/api/export/overview.csv', rows).headers['X-Export-Rows'], '4')

    def test_empty_export_keeps_headers_without_fabricated_data(self):
        response = self.get('/api/export/temporal.csv', [])
        self.assertEqual(len(list(csv.reader(StringIO(response.data.decode('utf-8-sig'))))), 1)
        response = self.get('/api/export/temporal.xlsx', [])
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual(workbook['Traffic data'].max_row, 1)
        self.assertEqual(workbook['Traffic data'].max_column, len(COLUMNS))
        self.assertEqual(dict(workbook['Export info'].values)['จำนวนรถรวม (คัน)'], 0)
        workbook.close()


if __name__ == '__main__':
    unittest.main()
