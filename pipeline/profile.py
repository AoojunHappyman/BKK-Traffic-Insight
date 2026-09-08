"""Inspect raw workbook layout without assuming a header or traffic schema."""

import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import logging
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LOGGER = logging.getLogger(__name__)


def fingerprint(path):
    digest = hashlib.sha256()
    with path.open('rb') as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def summarize_rows(rows):
    """Blank cells include layout padding; duplicates are not business-key checks."""
    width = max((len(row) for row in rows), default=0)
    nonempty = [row for row in rows if any(value is not None and value != '' for value in row)]
    columns = []
    for index in range(width):
        values = [row[index] if index < len(row) else None for row in rows]
        present = [value for value in values if value is not None and value != '']
        columns.append({
            'column': get_column_letter(index + 1),
            'blank_cells': len(values) - len(present),
            'observed_python_types': dict(Counter(type(value).__name__ for value in present)),
        })
    # Include types to distinguish numeric cells from strings with the same text.
    signatures = [tuple((type(value).__name__, str(value)) for value in row) for row in nonempty]
    return {
        'rows_read': len(rows),
        'columns_read': width,
        'nonempty_rows': len(nonempty),
        'duplicate_nonempty_rows_after_first': len(signatures) - len(set(signatures)),
        'columns': columns,
        'preview': [
            {'row': number, 'cells': {get_column_letter(i + 1): value
                                     for i, value in enumerate(row) if value is not None and value != ''}}
            for number, row in enumerate(rows[:14], start=1)
        ],
    }


def inspect_file(path, encoding='utf-8-sig', delimiter=','):
    path = Path(path).resolve()
    if path.suffix.lower() not in {'.xlsx', '.csv'}:
        raise ValueError('Supported input formats: .xlsx and .csv')
    checksum = fingerprint(path)
    result = {'file': path.name, 'source_path': str(path), 'sha256': checksum,
              'size_bytes': path.stat().st_size, 'sheets': []}
    if path.suffix.lower() == '.xlsx':
        # Normal read mode exposes merged ranges; the workbook is never saved.
        workbook = load_workbook(path, data_only=False, keep_links=False)
        try:
            for sheet in workbook:
                rows = list(sheet.iter_rows(values_only=True))
                summary = summarize_rows(rows)
                summary.update({
                    'name': sheet.title,
                    'visibility': sheet.sheet_state,
                    'merged_ranges': [str(item) for item in sheet.merged_cells.ranges],
                    'formula_cells': [cell.coordinate for row in sheet for cell in row if cell.data_type == 'f'],
                    'excel_error_cells': [cell.coordinate for row in sheet for cell in row if cell.data_type == 'e'],
                })
                result['sheets'].append(summary)
        finally:
            workbook.close()
    else:
        # Keep header rows and literal NA/NULL strings; no automatic date conversion.
        try:
            frame = pd.read_csv(path, header=None, dtype=str, keep_default_na=False,
                                encoding=encoding, sep=delimiter, skip_blank_lines=False)
            rows = frame.values.tolist()
        except pd.errors.EmptyDataError:
            rows = []
        result['sheets'].append({'name': 'csv', **summarize_rows(rows)})
    if fingerprint(path) != checksum:
        raise RuntimeError(f'Source changed while being inspected: {path.name}')
    return result


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path, help='CSV/XLSX file or directory (non-recursive)')
    parser.add_argument('--output', type=Path, default=Path('reports/inventory.json'))
    parser.add_argument('--encoding', default='utf-8-sig', help='CSV encoding; e.g. cp874')
    parser.add_argument('--delimiter', default=',', help='Single-character CSV delimiter')
    args = parser.parse_args(argv)
    if not args.source.exists():
        parser.error('Source does not exist')
    if args.output.suffix.lower() != '.json':
        parser.error('Output must be a .json file')
    if len(args.delimiter) != 1:
        parser.error('Delimiter must be one character')
    if args.source.is_dir():
        paths = sorted(p for p in args.source.iterdir()
                       if p.is_file() and p.suffix.lower() in {'.xlsx', '.csv'} and not p.name.startswith('~$'))
    else:
        paths = [args.source]
    if not paths:
        parser.error('No CSV/XLSX files found')
    if any(args.output.resolve() == path.resolve() for path in paths):
        parser.error('Output cannot overwrite an input')
    report = {'generated_at': datetime.now(timezone.utc).isoformat(),
              'scope': 'Raw layout inspection only; no cleaning, date inference or database loading.',
              'files': [], 'errors': []}
    for path in paths:
        try:
            report['files'].append(inspect_file(path, args.encoding, args.delimiter))
            LOGGER.info('Inspected %s', path.name)
        except Exception as error:
            LOGGER.error('Could not inspect %s: %s', path.name, error)
            report['errors'].append({'file': path.name, 'error': str(error)})
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    LOGGER.info('Report: %s (%s files, %s errors)', args.output, len(report['files']), len(report['errors']))
    return 1 if report['errors'] else 0


if __name__ == '__main__':
    logging.basicConfig(level=logging.INFO, format='%(levelname)s %(message)s')
    raise SystemExit(main())
