"""Clean the inspected BKK intersection report layout, retaining audit evidence.

Run: python -m pipeline.clean SOURCE_DIRECTORY --output data/processed/RUN_NAME
Only literal, nonnegative integer vehicle counts are accepted. Entire survey
blocks are quarantined on errors so downstream totals cannot look complete when
one road or interval was rejected. No Excel source is ever saved.
"""

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from pipeline.profile import fingerprint

VERSION = '1.0'
FULL_MONTHS = 'มกราคม กุมภาพันธ์ มีนาคม เมษายน พฤษภาคม มิถุนายน กรกฎาคม สิงหาคม กันยายน ตุลาคม พฤศจิกายน ธันวาคม'.split()
SHORT_MONTHS = 'ม.ค. ก.พ. มี.ค. เม.ย. พ.ค. มิ.ย. ก.ค. ส.ค. ก.ย. ต.ค. พ.ย. ธ.ค.'.split()
DATE_PATTERN = re.compile(r'\(?\s*(?<!\d)\d{1,2}\s*(?:' + '|'.join(map(re.escape, SHORT_MONTHS)) + r')\s*(?:\d{4}|\d{2})(?!\d)\s*\)?')
CATEGORIES = ['passenger_car', 'van_pickup', 'large_bus', 'small_bus', 'truck', 'three_wheeler']
THAI_CATEGORIES = ['รถยนต์นั่ง', 'ตู้/ปิคอัพ', 'เมล์ใหญ่', 'เมล์เล็ก', 'บรรทุก', 'สามล้อ']


def text(value):
    return ' '.join(str(value).split()) if value is not None else ''


def stable_id(*parts):
    return hashlib.sha256(json.dumps(parts, ensure_ascii=False).encode('utf-8')).hexdigest()


def report_month(value):
    match = re.search(r'ประจำเดือน\s*(' + '|'.join(FULL_MONTHS) + r')\s*(\d{4})', text(value))
    if not match:
        raise ValueError('Unrecognized report month in A1')
    year = int(match[2]) - 543
    return date(year, FULL_MONTHS.index(match[1]) + 1, 1)


def survey_date(value, month):
    match = re.fullmatch(r'(\d{1,2})\s*([^\s\d]+)\s*(\d{2}|\d{4})', text(value).strip('() '))
    if not match or match[2] not in SHORT_MONTHS:
        raise ValueError('Unrecognized survey date')
    year = int(match[3])
    if len(match[3]) == 2:
        year += ((month.year + 543) // 100) * 100
    result = date(year - 543, SHORT_MONTHS.index(match[2]) + 1, int(match[1]))
    if (result.year, result.month) != (month.year, month.month):
        raise ValueError('Survey date does not match report month')
    return result


def period(value):
    match = re.fullmatch(r'.+\((\d{1,2})\.(\d{2})\s*-\s*(\d{1,2})\.(\d{2})\s*น\.\)', text(value))
    if not match:
        raise ValueError('Unrecognized observation period')
    h1, m1, h2, m2 = map(int, match.groups())
    if not (0 <= h1 <= 23 and 0 <= h2 <= 23 and 0 <= m1 <= 59 and 0 <= m2 <= 59):
        raise ValueError('Invalid observation time')
    duration = (h2 * 60 + m2) - (h1 * 60 + m1)
    if duration <= 0:
        raise ValueError('Observation period must end after start')
    return f'{h1:02}:{m1:02}:00', f'{h2:02}:{m2:02}:00', duration


def count(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError('Count must be a literal number, not a formula or missing value')
    if not math.isfinite(value) or value < 0 or value != int(value):
        raise ValueError('Count must be a finite nonnegative integer')
    return int(value)


def coordinates(value):
    match = re.fullmatch(r'\s*([-+]?\d+(?:\.\d+)?)\s*[,\s]\s*([-+]?\d+(?:\.\d+)?)\s*', str(value))
    if not match:
        raise ValueError('Coordinates must contain exactly latitude and longitude')
    lat, lon = map(float, match.groups())
    if not (-90 <= lat <= 90 and -180 <= lon <= 180):
        raise ValueError('Coordinates outside valid geographic ranges')
    return lat, lon


def evaluate_sum(sheet, address, seen=None):
    """Evaluate only same-sheet SUM(cell/range), never arbitrary Excel or Python."""
    seen = set() if seen is None else set(seen)
    if address in seen or len(seen) > 20:
        raise ValueError('Circular or excessively nested subtotal formula')
    seen.add(address)
    value = sheet[address].value
    if value is None:
        return 0  # Blank cells inside a subtotal range are layout padding.
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return count(value)
    match = re.fullmatch(r'=SUM\(([A-Z]+\d+(?::[A-Z]+\d+)?)\)', str(value), re.I)
    if not match:
        raise ValueError('Unsupported formula or nonnumeric subtotal')
    c1, r1, c2, r2 = range_boundaries(match[1].upper())
    if c2 > sheet.max_column or r2 > sheet.max_row or (c2 - c1 + 1) * (r2 - r1 + 1) > 10000:
        raise ValueError('Subtotal range outside inspected sheet bounds')
    return sum(evaluate_sum(sheet, cell.coordinate, seen)
               for row in sheet.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2) for cell in row)


def clean_sheet(sheet, cached, source_id):
    month = report_month(sheet['A1'].value)
    if [text(sheet.cell(4, c).value) for c in range(5, 11)] != THAI_CATEGORIES:
        raise ValueError('Vehicle category headers differ from the supported layout')
    if [text(sheet.cell(3, c).value) for c in (1, 3, 4)] != ['ลำดับที่', 'ถนน/ซอย', 'ช่วงเวลา']:
        raise ValueError('Main headers differ from the supported layout')
    if [text(sheet.cell(4, c).value) for c in (11, 12, 13)] != ['แต่ละช่วงเวลา', 'แต่ละถนน', 'รวมทั้งแยก']:
        raise ValueError('Subtotal headers differ from the supported layout')
    starts = [r for r in range(5, sheet.max_row + 1) if sheet.cell(r, 1).value is not None]
    if not starts:
        raise ValueError('No survey blocks found')
    if any(sheet.cell(r, c).value is not None for r in range(5, starts[0]) for c in range(2, 14)):
        raise ValueError('Unassigned data before the first survey block')
    surveys, roads, observations, issues, raw_rows = [], [], [], [], []
    for begin, end in zip(starts, starts[1:] + [sheet.max_row + 1]):
        sid = stable_id(source_id, sheet.title, begin)
        survey = {'survey_id': sid, 'source_id': source_id, 'sheet_name': sheet.title,
                  'source_start_row': begin, 'source_end_row': end - 1,
                  'source_sequence': sheet.cell(begin, 1).value, 'report_month': month.isoformat(),
                  'survey_date': None, 'survey_date_raw': None, 'survey_date_source_row': None,
                  'intersection_name_raw': None, 'intersection_name': None,
                  'coordinate_raw': None, 'latitude': None, 'longitude': None}

        def issue(code, detail, row=begin, severity='error'):
            issues.append({'source_id': source_id, 'survey_id': sid, 'sheet_name': sheet.title,
                           'source_row': row, 'severity': severity, 'code': code, 'detail': detail})

        try:
            if count(survey['source_sequence']) < 1:
                raise ValueError('Survey sequence must be positive')
        except ValueError as error:
            issue('invalid_survey_sequence', str(error))
        for r in range(begin, end):
            extra = [sheet.cell(r, c).coordinate for c in range(15, sheet.max_column + 1)
                     if sheet.cell(r, c).value is not None]
            if extra:
                issue('extra_cells_outside_table', f'Not used as traffic facts: {extra}', r, 'warning')

        names, dates, coords = [], [], []
        for r in range(begin, end):
            b = sheet.cell(r, 2).value
            if b is not None:
                # Newer reports combine the name and date in one multiline cell.
                tokens = DATE_PATTERN.findall(str(b))
                dates.extend((r, token) for token in tokens)
                remaining = DATE_PATTERN.sub('', str(b)).strip()
                if remaining:
                    names.append((r, remaining))
            if sheet.cell(r, 14).value is not None:
                coords.append((r, sheet.cell(r, 14).value))
        if not names:
            issue('missing_name', 'No location name in this survey block')
        else:
            # Preserve all label lines, including direction/segment descriptions.
            survey['intersection_name_raw'] = '\n'.join(str(value) for _, value in names)
            survey['intersection_name'] = text(survey['intersection_name_raw'])
        if len(dates) != 1:
            issue('ambiguous_date', f'Expected one date, found {dates}')
        else:
            survey['survey_date_raw'] = str(dates[0][1])
            survey['survey_date_source_row'] = dates[0][0]
            try:
                survey['survey_date'] = survey_date(dates[0][1], month).isoformat()
            except ValueError as error:
                issue('invalid_date', str(error), dates[0][0])
        if len(coords) != 1:
            issue('missing_or_ambiguous_coordinates', f'Found {len(coords)} coordinate cells', severity='warning')
        else:
            survey['coordinate_raw'] = str(coords[0][1])
            try:
                survey['latitude'], survey['longitude'] = coordinates(coords[0][1])
                # Broad review envelope, not a Bangkok administrative boundary test.
                if not (13 <= survey['latitude'] <= 15 and 99 <= survey['longitude'] <= 102):
                    issue('coordinates_need_review', 'Outside broad Bangkok-area review envelope', severity='warning')
            except ValueError as error:
                issue('invalid_coordinates', str(error), coords[0][0], severity='warning')

        local_observations, local_roads = [], []
        road_starts = [begin]
        for r in range(begin + 1, end):
            try:
                if period(sheet.cell(r, 4).value)[0] == '07:00:00':
                    road_starts.append(r)
            except ValueError:
                pass  # Logged by observation validation below.
        road_at_row = {}
        for road_begin, road_end in zip(road_starts, road_starts[1:] + [end]):
            labels = [(r, sheet.cell(r, 3).value) for r in range(road_begin, road_end)
                      if sheet.cell(r, 3).value is not None]
            if len(labels) != 1:
                issue('missing_or_ambiguous_road', f'Expected one road label in interval group, found {labels}', road_begin)
                continue
            road = {'road_id': stable_id(sid, road_begin), 'survey_id': sid,
                    'source_start_row': road_begin, 'road_name_source_row': labels[0][0],
                    'road_name_raw': str(labels[0][1]), 'road_name': text(labels[0][1])}
            local_roads.append(road)
            road_at_row.update({r: road for r in range(road_begin, road_end)})
        for r in range(begin, end):
            values = [sheet.cell(r, c).value for c in range(1, 15)]
            raw_rows.append({'survey_id': sid, 'source_id': source_id, 'sheet_name': sheet.title,
                             'source_row': r, 'cells': values,
                             'cached_cells': [cached.cell(r, c).value for c in range(1, 15)]})
            active_road = road_at_row.get(r)
            if values[3] is None:
                if any(v is not None for v in values[4:13]):
                    issue('unassigned_numeric_row', 'Numeric/subtotal row has no observation period', r)
                continue
            if active_road is None:
                issue('missing_road', 'No road header inside this survey block', r)
                continue
            try:
                start, finish, duration = period(values[3])
                counts = [count(v) for v in values[4:10]]
            except ValueError as error:
                issue('invalid_observation', str(error), r)
                continue
            record = {'observation_id': stable_id(source_id, sheet.title, r),
                      'survey_id': sid, 'road_id': active_road['road_id'], 'source_row': r,
                      'period_raw': str(values[3]), 'period_start': start, 'period_end': finish,
                      'duration_minutes': duration, **dict(zip(CATEGORIES, counts)),
                      'vehicle_total': sum(counts)}
            local_observations.append(record)

        def reconcile(address, expected):
            try:
                if sheet[address].value is None:
                    raise ValueError('Missing subtotal')
                calculated = evaluate_sum(sheet, address)
                if calculated != expected:
                    raise ValueError(f'Expected {expected}, recalculated {calculated}')
                cache = cached[address].value
                if cache is None:
                    issue('missing_formula_cache', address, sheet[address].row, 'warning')
                elif cache != calculated:
                    issue('stale_formula_cache', f'{address}: cached={cache}, recalculated={calculated}',
                          sheet[address].row, 'warning')
            except ValueError as error:
                issue('subtotal_mismatch_or_unverifiable', f'{address}: {error}', sheet[address].row)

        for record in local_observations:
            reconcile(f"K{record['source_row']}", record['vehicle_total'])
        for index, road in enumerate(local_roads):
            road_end = local_roads[index + 1]['source_start_row'] if index + 1 < len(local_roads) else end
            rows = [o for o in local_observations if o['road_id'] == road['road_id']]
            windows = sorted((o['period_start'], o['period_end']) for o in rows)
            if len(windows) != 3 or not windows or windows[0][0] != '07:00:00' or windows[-1][1] != '19:00:00' or any(a[1] != b[0] for a, b in zip(windows, windows[1:])):
                issue('incomplete_or_overlapping_periods', 'Expected three contiguous intervals covering 07:00–19:00', road['source_start_row'])
            totals = [r for r in range(road['source_start_row'], road_end) if sheet.cell(r, 12).value is not None]
            if len(totals) != 1:
                issue('ambiguous_road_total', f'Found {len(totals)} road totals', road['source_start_row'])
            else:
                reconcile(f'L{totals[0]}', sum(o['vehicle_total'] for o in rows))
        totals = [r for r in range(begin, end) if sheet.cell(r, 13).value is not None]
        if len(totals) != 1:
            issue('ambiguous_survey_total', f'Found {len(totals)} survey totals')
        else:
            reconcile(f'M{totals[0]}', sum(o['vehicle_total'] for o in local_observations))
        if not local_observations:
            issue('empty_survey', 'No valid observations')
        surveys.append(survey)
        roads.extend(local_roads)
        observations.extend(local_observations)
    return surveys, roads, observations, issues, raw_rows


def run(source, output):
    source, output = Path(source).resolve(), Path(output).resolve()
    paths = sorted(p for p in source.glob('*.xlsx') if not p.name.startswith('~$'))
    if not paths:
        raise ValueError('No XLSX files found in source directory')
    if output == source or source in output.parents:
        raise ValueError('Output must be outside the raw source directory')
    if output.exists():
        raise ValueError('Output directory already exists; use a new run name to preserve audit history')
    sources, surveys, roads, observations, issues, raw_rows = [], [], [], [], [], []
    for path in paths:
        checksum = fingerprint(path)
        source_id = stable_id(path.name, checksum)
        sources.append({'source_id': source_id, 'filename': path.name, 'sha256': checksum, 'source_path': str(path)})
        try:
            workbook = load_workbook(path, data_only=False, keep_links=False)
            try:
                cached = load_workbook(path, data_only=True, keep_links=False)
                try:
                    for sheet in workbook:
                        result = clean_sheet(sheet, cached[sheet.title], source_id)
                        for target, values in zip((surveys, roads, observations, issues, raw_rows), result):
                            target.extend(values)
                finally:
                    cached.close()
            finally:
                workbook.close()
        except Exception as error:
            issues.append({'source_id': source_id, 'survey_id': None, 'sheet_name': None,
                           'source_row': None, 'severity': 'error', 'code': 'file_error', 'detail': str(error)})
        if fingerprint(path) != checksum:
            raise RuntimeError(f'Source changed during run: {path.name}')

    survey_lookup = {s['survey_id']: s for s in surveys}
    road_lookup = {r['road_id']: r for r in roads}
    business_keys = defaultdict(list)
    for record in observations:
        s, road = survey_lookup[record['survey_id']], road_lookup[record['road_id']]
        # Candidate duplicates only: never merge names, roads or survey identities.
        if s['survey_date'] and s['intersection_name']:
            key = (s['survey_date'], s['intersection_name'], road['road_name'], record['period_start'], record['period_end'])
            business_keys[key].append(record)
    for records in business_keys.values():
        if len(records) > 1:
            for record in records:
                s = survey_lookup[record['survey_id']]
                issues.append({'source_id': s['source_id'], 'survey_id': s['survey_id'], 'sheet_name': s['sheet_name'],
                               'source_row': record['source_row'], 'severity': 'error', 'code': 'duplicate_candidate',
                               'detail': 'Repeated date/name/road/period key; all affected surveys retained for review'})
    failed_sources = {i['source_id'] for i in issues if i['code'] == 'file_error'}
    blocked = {i['survey_id'] for i in issues if i['severity'] == 'error' and i['survey_id']}
    blocked.update(s['survey_id'] for s in surveys if s['source_id'] in failed_sources)
    for survey in surveys:
        survey['quality_status'] = 'quarantined' if survey['survey_id'] in blocked else 'accepted'
    accepted_surveys = [s for s in surveys if s['survey_id'] not in blocked]
    accepted_roads = [r for r in roads if r['survey_id'] not in blocked]
    accepted_observations = [o for o in observations if o['survey_id'] not in blocked]
    summary = {'pipeline_version': VERSION, 'created_at': datetime.now(timezone.utc).isoformat(),
               'source_files': len(sources), 'surveys_total': len(surveys), 'surveys_accepted': len(accepted_surveys),
               'surveys_quarantined': len(blocked), 'observations_accepted': len(accepted_observations),
               'observation_rows_seen': sum(r['cells'][3] is not None for r in raw_rows),
               'observation_rows_quarantined': sum(r['cells'][3] is not None and r['survey_id'] in blocked for r in raw_rows),
               'issues_by_code': dict(Counter(i['code'] for i in issues)),
               'error_count': sum(i['severity'] == 'error' for i in issues),
               'warning_count': sum(i['severity'] == 'warning' for i in issues),
               'notes': 'Accepted exports exclude entire quarantined surveys. Counts are surveyed intervals, not continuous monthly traffic.'}
    if summary['observation_rows_seen'] != summary['observations_accepted'] + summary['observation_rows_quarantined']:
        raise RuntimeError('Observation row accounting failed; no output was written')
    output.mkdir(parents=True)
    for name, values in [('sources', sources), ('surveys', accepted_surveys), ('roads', accepted_roads),
                         ('observations', accepted_observations)]:
        # JSON is canonical; CSV is a convenient derived view. No Excel workbooks are authored.
        (output / f'{name}.json').write_text(json.dumps(values, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
        if values:
            pd.DataFrame(values).to_csv(output / f'{name}.csv', index=False, encoding='utf-8-sig')
    for name, values in [('summary', summary), ('issues', issues), ('all_surveys', surveys),
                         ('rejected_rows', [r for r in raw_rows if r['survey_id'] in blocked])]:
        (output / f'{name}.json').write_text(json.dumps(values, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    return summary


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--output', type=Path, required=True, help='New directory for this run')
    args = parser.parse_args(argv)
    try:
        summary = run(args.source, args.output)
    except (ValueError, OSError, RuntimeError) as error:
        parser.exit(1, f'Cleaning failed: {error}\n')
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 2 if summary['error_count'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
