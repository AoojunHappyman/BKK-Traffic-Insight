"""Read original XLSX evidence for quarantined surveys; never replace source values.

python -m pipeline.audit_quarantine SOURCE --baseline CLEAN_RUN --output NEW_DIRECTORY
The JSON evidence includes cell addresses, literal/formula values, saved caches,
external-link snapshots, source hashes and the existing cleaner's fresh verdict.
Caches and SUMs of empty cells are diagnostic evidence, not verified corrections.
"""

import argparse
from collections import Counter
from datetime import datetime, timezone
import json
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from pipeline.clean import clean_sheet, count, evaluate_sum
from pipeline.profile import fingerprint

NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
EXTERNAL = re.compile(r"=\[(\d+)\]([^!]+)!\$?([A-Z]+)\$?(\d+)$")


def external_snapshots(path):
    """Inspect embedded link caches only. Never follow or refresh external URLs."""
    result = {}
    with ZipFile(path) as archive:
        for member in archive.namelist():
            match = re.fullmatch(r'xl/externalLinks/externalLink(\d+)\.xml', member)
            if not match:
                continue
            root = ET.fromstring(archive.read(member))
            names = [s.attrib['val'] for s in root.findall('.//s:sheetName', NS)]
            cells = {}
            for sheet in root.findall('.//s:sheetData', NS):
                name = names[int(sheet.attrib['sheetId'])]
                for cell in sheet.findall('.//s:cell', NS):
                    cells[f"{name}!{cell.attrib['r']}"] = {
                        'type': cell.get('t', 'n'), 'value': cell.findtext('s:v', namespaces=NS)}
            relationships = member.replace('xl/externalLinks/', 'xl/externalLinks/_rels/') + '.rels'
            targets = []
            if relationships in archive.namelist():
                targets = [r.attrib['Target'] for r in ET.fromstring(archive.read(relationships))]
            result[match[1]] = {'archive_member': member, 'targets': targets,
                               'saved_cells': cells,
                               'limitation': 'Saved external values, not the upstream workbook or its formulas.'}
    return result


def cell_evidence(sheet, cached, row, column, snapshots):
    cell = sheet.cell(row, column)
    evidence = {'address': cell.coordinate, 'value': cell.value,
                'cached_value': cached.cell(row, column).value, 'type': cell.data_type,
                'number_format': cell.number_format}
    if column in range(5, 11):
        try:
            count(cell.value)
            evidence['count_status'] = 'valid_literal'
        except ValueError:
            evidence['count_status'] = ('missing' if cell.value is None else
                                        'formula' if cell.data_type == 'f' else 'invalid_literal')
    match = EXTERNAL.fullmatch(str(cell.value))
    if match:
        link, name, col, r = match.groups()
        evidence['external_reference'] = {
            'link_id': link, 'cell': f'{name}!{col}{r}',
            'saved_external_cell': snapshots.get(link, {}).get('saved_cells', {}).get(f'{name}!{col}{r}'),
            'upstream_verified': False,
        }
    if column in (11, 12, 13) and cell.value is not None:
        try:
            evidence['recalculated_from_literal_counts'] = evaluate_sum(sheet, cell.coordinate)
        except ValueError as error:
            evidence['recalculation_error'] = str(error)
    return evidence


def audit(source, baseline, output):
    source, baseline, output = (Path(p).resolve() for p in (source, baseline, output))
    if output.exists() or output == source or source in output.parents:
        raise ValueError('Use a new output directory outside the original source directory')

    def read(name):
        return json.loads((baseline / f'{name}.json').read_text(encoding='utf-8'))

    sources = {s['source_id']: s for s in read('sources')}
    quarantined = [s for s in read('all_surveys') if s['quality_status'] == 'quarantined']
    baseline_raw = read('rejected_rows')
    baseline_issues = read('issues')
    groups, manifests = [], []
    for source_id in sorted({s['source_id'] for s in quarantined}):
        original = sources[source_id]
        filename = original['filename']
        if Path(filename).name != filename:
            raise ValueError('Source filename must not contain directories')
        path = source / filename
        before = fingerprint(path)
        if before != original['sha256']:
            raise ValueError(f'Source revision differs from baseline: {filename}. Review the new revision separately.')
        snapshots = external_snapshots(path)
        book = load_workbook(path, data_only=False, keep_links=False)
        cached_book = load_workbook(path, data_only=True, keep_links=False)
        try:
            for sheet_name in sorted({s['sheet_name'] for s in quarantined if s['source_id'] == source_id}):
                sheet, cached = book[sheet_name], cached_book[sheet_name]
                fresh = clean_sheet(sheet, cached, source_id)
                for survey in (s for s in quarantined if s['source_id'] == source_id and s['sheet_name'] == sheet_name):
                    sid = survey['survey_id']
                    raw = [r for r in fresh[4] if r['survey_id'] == sid]
                    expected_raw = [r for r in baseline_raw if r['survey_id'] == sid]
                    if raw != expected_raw:
                        raise ValueError(f'Raw/cached evidence differs from baseline for {filename}:{sid}')
                    issues = [i for i in fresh[3] if i['survey_id'] == sid]
                    previous = [i for i in baseline_issues if i['survey_id'] == sid]
                    same_issues = sorted(json.dumps(i, sort_keys=True) for i in issues) == sorted(json.dumps(i, sort_keys=True) for i in previous)
                    begin, end = survey['source_start_row'], survey['source_end_row']
                    cells = [cell_evidence(sheet, cached, r, c, snapshots)
                             for r in range(begin, end + 1) for c in range(1, 15)]
                    status = Counter(c['count_status'] for c in cells if 'count_status' in c)
                    errors = [i for i in issues if i['severity'] == 'error']
                    # This is an evidence collector. A changed validation result still
                    # needs explicit source review before any correction/import.
                    groups.append({'survey': survey, 'filename': filename,
                                   'source_sha256': before, 'range': f'A{begin}:N{end}',
                                   'count_cells': dict(status), 'cells': cells,
                                   'merged_ranges': sorted(str(m) for m in sheet.merged_cells.ranges
                                                           if m.min_row <= end and m.max_row >= begin),
                                   'fresh_issues': issues, 'issues_match_baseline': same_issues,
                                   'fresh_error_count': len(errors),
                                   'decision': 'remain_quarantined' if errors else 'requires_source_review',
                                   'proven_corrections': [], 'external_snapshots': snapshots})
        finally:
            book.close()
            cached_book.close()
        after = fingerprint(path)
        if before != after:
            raise RuntimeError(f'Source changed during audit: {filename}')
        manifests.append({'filename': filename, 'source_id': source_id,
                          'sha256_before': before, 'sha256_after': after})
    report = {'created_at': datetime.now(timezone.utc).isoformat(),
              'source_directory': str(source), 'baseline_directory': str(baseline),
              'policy': 'No rounding, zero filling, deleting unlabeled periods or accepting external caches as verified counts.',
              'source_files': manifests, 'groups': groups,
              'summary': {'groups_reviewed': len(groups),
                          'observation_rows_reviewed': len(baseline_raw),
                          'groups_still_quarantined': sum(g['decision'] == 'remain_quarantined' for g in groups),
                          'count_cells': dict(sum((Counter(g['count_cells']) for g in groups), Counter())),
                          'proven_corrections': 0, 'groups_released': 0}}
    output.mkdir(parents=True)
    (output / 'evidence.json').write_text(json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    return report['summary']


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--baseline', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(audit(args.source, args.baseline, args.output), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
