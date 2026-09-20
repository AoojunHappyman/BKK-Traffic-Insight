"""Download accepted observations; spreadsheet strings are never executable formulas."""
import csv
from datetime import date, datetime, timezone
from io import BytesIO, StringIO

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

COLUMNS = [
    ('survey_date', 'วันที่สำรวจ'), ('intersection_name', 'สถานที่'), ('road_name', 'ถนน'),
    ('period_start', 'เริ่มช่วงสำรวจ'), ('period_end', 'สิ้นสุดช่วงสำรวจ'),
    ('duration_minutes', 'ระยะเวลาสำรวจ (นาที)'),
    ('passenger_car', 'รถยนต์นั่ง (คัน)'), ('van_pickup', 'ตู้/ปิคอัพ (คัน)'),
    ('large_bus', 'เมล์ใหญ่ (คัน)'), ('small_bus', 'เมล์เล็ก (คัน)'),
    ('truck', 'บรรทุก (คัน)'), ('three_wheeler', 'สามล้อ (คัน)'),
    ('vehicle_total', 'รวมรถ (คัน)'), ('latitude', 'ละติจูด'), ('longitude', 'ลองจิจูด'),
    ('filename', 'ไฟล์ต้นฉบับ'), ('sheet_name', 'ชีตต้นฉบับ'), ('source_row', 'แถวต้นฉบับ'),
    ('survey_id', 'รหัสกลุ่มสำรวจ'), ('observation_id', 'รหัสแถวข้อมูล'),
]


def csv_value(value):
    # Quoting alone does not stop spreadsheet formula injection.
    if isinstance(value, str) and (value.lstrip().startswith(('=', '+', '-', '@')) or
                                   value.startswith(('\t', '\r', '\n'))):
        return "'" + value
    return value


def download(rows, view, format_name, applied_filters):
    headers = [label for _, label in COLUMNS]
    timestamp = datetime.now(timezone.utc)
    filename = f'bangkok-traffic-{view}-{timestamp:%Y%m%d-%H%M%S}.{format_name}'
    if format_name == 'csv':
        text = StringIO(newline='')
        writer = csv.writer(text)
        writer.writerow(headers)
        writer.writerows([csv_value(row.get(key)) for key, _ in COLUMNS] for row in rows)
        content = BytesIO(text.getvalue().encode('utf-8-sig'))
        mimetype = 'text/csv; charset=utf-8'
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Traffic data'
        sheet.append(headers)
        for row_number, row in enumerate(rows, 2):
            values = [row.get(key) for key, _ in COLUMNS]
            if values[0]:
                values[0] = date.fromisoformat(values[0])
            sheet.append(values)
            for column in range(1, len(COLUMNS) + 1):
                cell = sheet.cell(row_number, column)
                if isinstance(cell.value, str):
                    cell.data_type = 's'
            sheet.cell(row_number, 1).number_format = 'yyyy-mm-dd'
        sheet.freeze_panes = 'D2'
        sheet.auto_filter.ref = sheet.dimensions
        for index, (key, _) in enumerate(COLUMNS, 1):
            sheet.column_dimensions[get_column_letter(index)].width = (
                34 if key in {'intersection_name', 'road_name', 'filename'} else
                22 if key.endswith('_id') else 20)
        notes = workbook.create_sheet('Export info')
        notes.append(['รายการ', 'รายละเอียด'])
        info = [
            ('หน้า', view), ('เวลาส่งออก (UTC)', timestamp.isoformat()),
            ('จำนวนแถวช่วงเวลา', len(rows)),
            ('จำนวนกลุ่มสำรวจ', len({r['survey_id'] for r in rows})),
            ('จำนวนรถรวม (คัน)', sum(r['vehicle_total'] for r in rows)),
            ('วันที่เริ่มต้น', applied_filters.get('start_date', 'ไม่จำกัด')),
            ('วันที่สิ้นสุด', applied_filters.get('end_date', 'ไม่จำกัด')),
            ('สถานที่', applied_filters.get('intersection_name', 'ทุกสถานที่')),
            ('กลุ่มสำรวจ', applied_filters.get('survey_id', 'ทุกกลุ่ม')),
            ('ช่วงเวลา', applied_filters.get('period', 'ทุกช่วงเวลา')),
            ('หน่วยของแต่ละแถว', 'ถนน × วันสำรวจ × ช่วงเวลา; ยอดรถเป็นคัน'),
            ('ขอบเขต', 'ข้อมูลที่ผ่านการตรวจเท่านั้น ไม่รวมกลุ่มที่กักไว้และปี 2022'),
            ('ข้อจำกัด', 'ข้อมูลสำรวจ ไม่ใช่ข้อมูลสดหรือยอดรายเดือนต่อเนื่อง ช่วงเวลามีความยาวต่างกัน'),
            ('วันหยุด', 'หากแบ่งประเภทวัน ใช้เสาร์–อาทิตย์ ไม่รวมปฏิทินวันหยุดราชการ'),
            ('พิกัดในหน้า Map', 'เฉพาะกลุ่มที่มีพิกัดพร้อมใช้งาน ไม่รวมจุดที่รอตรวจพิกัด' if view == 'map'
             else 'ช่องว่างหมายถึงไม่มีพิกัด กลุ่มที่รอตรวจพิกัดยังมีปริมาณรถในชุดวิเคราะห์'),
        ]
        for pair in info:
            notes.append(pair)
            for cell in notes[notes.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
        notes.column_dimensions['A'].width = 28
        notes.column_dimensions['B'].width = 100
        notes.freeze_panes = 'A2'
        for page in workbook:
            for cell in page[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='304A42')
            for row in page.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical='top')
        content = BytesIO()
        workbook.save(content)
        workbook.close()
        content.seek(0)
        mimetype = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response = send_file(content, mimetype=mimetype, as_attachment=True, download_name=filename, max_age=0)
    response.headers['Cache-Control'] = 'no-store'
    response.headers['X-Export-Rows'] = str(len(rows))
    return response
